"""
财务对账工具 - FastAPI 入口
"""

import hashlib
import io
import os
import re
import secrets
import time
from typing import Optional
from urllib.parse import quote

from fastapi import FastAPI, File, Header, HTTPException, UploadFile, Depends
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
import openpyxl

import database as db

app = FastAPI(title="财务对账工具")

# ── Session Token 管理（内存） ────────────────────────
# token → {"user_id": int, "username": str, "expires_at": float}
_active_sessions: dict[str, dict] = {}
SESSION_TTL = 8 * 3600  # 8 小时有效期


def _cleanup_expired():
    """清除过期 session"""
    now = time.time()
    expired = [t for t, s in _active_sessions.items() if s["expires_at"] < now]
    for t in expired:
        del _active_sessions[t]


# ── 鉴权依赖 ─────────────────────────────────────────

def require_auth(authorization: str = Header(None)):
    """验证 Bearer token，未通过则返回 401"""
    _cleanup_expired()
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(401, "未登录")
    token = authorization[7:]
    session = _active_sessions.get(token)
    if not session or session["expires_at"] < time.time():
        raise HTTPException(401, "登录已过期，请重新登录")
    return session


# ── 启动时初始化数据库 ───────────────────────────────
@app.on_event("startup")
def startup():
    db.init_db()


# ── Pydantic 请求模型 ────────────────────────────────

class LoginRequest(BaseModel):
    password: str


# ── URL 校验 ──────────────────────────────────────────

_URL_PATTERN = re.compile(r'^https?://\S+$', re.IGNORECASE)


def _validate_image_url(url: str) -> Optional[str]:
    """校验图片URL，合法返回 None，否则返回错误信息"""
    if not url or not url.strip():
        return None  # 选填，空值合法
    if not _URL_PATTERN.match(url.strip()):
        return "图片链接必须以 http:// 或 https:// 开头"
    return None


class ProductCreate(BaseModel):
    product_name: str = Field(default="", description="商品名称")
    online_sku: str = Field(default="", description="线上SKU（选填）")
    offline_sku: str = Field(..., min_length=1, description="线下SKU")
    size: str = Field(default="", description="尺码（选填）")
    cost: float = Field(..., ge=0, description="成本")
    price: float = Field(default=0, ge=0, description="销售价（选填）")
    colors: str = Field(default="", description="颜色，逗号分隔")
    image_url: str = Field(default="", description="图片链接（选填，仅支持 http/https URL）")


class ProductUpdate(BaseModel):
    product_name: str = Field(default="", description="商品名称")
    online_sku: str = Field(default="", description="线上SKU（选填）")
    offline_sku: str = Field(..., min_length=1)
    size: str = Field(default="")
    cost: float = Field(..., ge=0)
    price: float = Field(default=0, ge=0)
    colors: str = ""
    image_url: str = Field(default="", description="图片链接（选填，仅支持 http/https URL）")


class ClearAllRequest(BaseModel):
    password: str = Field(..., min_length=1, description="管理员密码")


# ── 货款对账模型 ──────────────────────────────────────

class PaymentItemSchema(BaseModel):
    offline_sku: str
    quantity: float = Field(..., description="数量，支持正负（负数为退货）")
    product_name: str = ""
    cost_price: float = 0
    cost_amount: float = 0


class PaymentCreate(BaseModel):
    name: str = Field(..., min_length=1, description="货账名")
    payment_date: str = Field(..., description="日期（yyyy-mm-dd）")
    items: list[PaymentItemSchema] = Field(..., min_items=1, description="明细列表")


class PaymentUpdate(BaseModel):
    name: str = Field(..., min_length=1)
    payment_date: str
    items: list[PaymentItemSchema] = Field(..., min_items=1)


class PaymentSettle(BaseModel):
    other_fee: float = Field(..., description="其他费用，允许0和负数")
    settlement_date: str = Field(..., description="结算日期（yyyy-mm-dd）")
    settlement_remark: str = Field("", description="备注")


# ── 登录 ─────────────────────────────────────────────

@app.post("/api/login")
def api_login(data: LoginRequest):
    user = db.verify_password(data.password)
    if not user:
        raise HTTPException(401, "密码错误")
    _cleanup_expired()
    token = secrets.token_hex(32)
    _active_sessions[token] = {
        "user_id": user["id"],
        "username": user["username"],
        "expires_at": time.time() + SESSION_TTL,
    }
    return {"token": token, "username": user["username"]}


@app.post("/api/logout")
def api_logout(authorization: str = Header(None)):
    if authorization and authorization.startswith("Bearer "):
        token = authorization[7:]
        _active_sessions.pop(token, None)
    return {"ok": True}


@app.get("/api/session")
def api_check_session(session: dict = Depends(require_auth)):
    """检查当前 session 是否有效"""
    return {"username": session["username"]}


# ── 商品 API ─────────────────────────────────────────

@app.get("/api/products")
def api_list_products(search: str = "", page: int = 1, page_size: int = 20,
                      _=Depends(require_auth)):
    return db.list_products(search, page, page_size)


@app.get("/api/products/{product_id}")
def api_get_product(product_id: int, _=Depends(require_auth)):
    p = db.get_product(product_id)
    if not p:
        raise HTTPException(404, "商品不存在")
    return p


@app.post("/api/products")
def api_create_product(data: ProductCreate, _=Depends(require_auth)):
    err = _validate_image_url(data.image_url)
    if err:
        raise HTTPException(400, err)
    try:
        return db.create_product(
            data.product_name, data.online_sku, data.offline_sku, data.size,
            data.cost, data.price, data.colors, data.image_url
        )
    except db.DuplicateOfflineSkuError as e:
        raise HTTPException(409, str(e))


@app.put("/api/products/{product_id}")
def api_update_product(product_id: int, data: ProductUpdate,
                       _=Depends(require_auth)):
    err = _validate_image_url(data.image_url)
    if err:
        raise HTTPException(400, err)
    try:
        p = db.update_product(
            product_id, data.product_name, data.online_sku, data.offline_sku, data.size,
            data.cost, data.price, data.colors, data.image_url
        )
    except db.DuplicateOfflineSkuError as e:
        raise HTTPException(409, str(e))
    if not p:
        raise HTTPException(404, "商品不存在")
    return p


@app.delete("/api/products/{product_id}")
def api_delete_product(product_id: int, _=Depends(require_auth)):
    ok = db.delete_product(product_id)
    if not ok:
        raise HTTPException(404, "商品不存在")
    return {"ok": True}


# ── 清空所有商品档案（需要登录 + 密码） ──────────────

# SHA256 of the clear-all password
_CLEAR_ALL_PASSWORD_HASH = "b6404bed39b3ba1b531531c94197b8982fbae8223fcff6309aeaa263fc46c08b"


@app.post("/api/products/clear")
def api_clear_all(data: ClearAllRequest, _=Depends(require_auth)):
    """清空所有商品档案，需验证管理员密码"""
    pw_hash = hashlib.sha256(data.password.encode()).hexdigest()
    if pw_hash != _CLEAR_ALL_PASSWORD_HASH:
        raise HTTPException(403, "密码错误，操作被拒绝")
    count = db.clear_all_products()
    return {"ok": True, "deleted": count}


# ── Excel 导出模板 / 批量导入 / 导出商品（需要登录） ──

EXPORT_HEADERS = ["商品名称", "线上SKU（选填）", "线下SKU（必填）", "尺码", "成本", "销售价", "颜色", "图片链接（选填）"]


def _build_export_xlsx(rows: list[dict], sheet_title: str) -> io.BytesIO:
    """根据数据行生成 xlsx 文件，返回 BytesIO 流"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(list(EXPORT_HEADERS))
    for p in rows:
        ws.append([
            p.get("product_name", ""),
            p.get("online_sku", ""),
            p.get("offline_sku", ""),
            p.get("size", ""),
            p.get("cost", 0),
            p.get("price", 0),
            p.get("colors", ""),
            p.get("image_url", ""),
        ])
    for col, w in zip("ABCDEFGH", [16, 18, 18, 14, 12, 12, 20, 28]):
        ws.column_dimensions[col].width = w
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


@app.get("/api/export-template")
def export_template(_=Depends(require_auth)):
    """下载空白导入模板（含示例数据）"""
    demo_rows = [
        {"product_name": "示例商品A", "online_sku": "SKU-001", "offline_sku": "OFF-001",
         "size": "1.5米", "cost": 50, "price": 80, "colors": "红色,蓝色",
         "image_url": "https://example.com/img.jpg"},
        {"product_name": "示例商品B", "online_sku": "", "offline_sku": "OFF-002",
         "size": "1.8米", "cost": 60, "price": 95, "colors": "灰色", "image_url": ""},
    ]
    stream = _build_export_xlsx(demo_rows, "商品导入模板")
    encoded = quote("商品导入模板.xlsx")
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"}
    )


@app.get("/api/export-products")
def export_products(search: str = "", _=Depends(require_auth)):
    """导出当前搜索结果的商品档案，格式与导入模板一致"""
    products = db.export_products(search)
    stream = _build_export_xlsx(products, "商品档案")
    # 带时间戳的文件名
    from datetime import datetime
    ts = datetime.now().strftime("%Y%m%d")
    filename = f"商品档案_{ts}.xlsx" if not search else f"商品档案_搜索结果_{ts}.xlsx"
    encoded = quote(filename)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"}
    )


@app.post("/api/import")
def api_import(file: UploadFile = File(...), _=Depends(require_auth)):
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "仅支持 .xlsx / .xls 文件")

    contents = file.file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents))
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(400, "Excel 至少需要表头 + 一行数据")

    header = [str(h).strip() if h else "" for h in rows[0]]
    # 标准化表头：移除括号注释（选填/必填），转为小写
    import re
    header_lower = [re.sub(r'[（(][^)）]*[)）]', '', h).strip().lower() for h in header]
    expected_old = ["商品名称", "线上sku", "线下sku", "尺码", "成本", "销售价", "颜色"]
    expected_new = expected_old + ["图片链接"]

    has_image_col = False
    if header_lower == expected_new:
        has_image_col = True
    elif header_lower == expected_old:
        has_image_col = False
    else:
        raise HTTPException(
            400,
            f"表头不匹配，期望: {', '.join(expected_old)}（或附带「图片链接（选填）」），实际: {', '.join(header_lower)}"
        )

    records = []
    errors = []
    for i, row in enumerate(rows[1:], start=2):
        if not row or all(v is None for v in row):
            continue
        try:
            image_url = ""
            if has_image_col:
                raw_url = str(row[7]).strip() if len(row) > 7 and row[7] else ""
                if raw_url:
                    err = _validate_image_url(raw_url)
                    if err:
                        errors.append(f"第{i}行: {err}（实际: {raw_url}）")
                        continue
                    image_url = raw_url
            rec = {
                "product_name": str(row[0]).strip() if row[0] else "",
                "online_sku": str(row[1]).strip() if row[1] else "",
                "offline_sku": str(row[2]).strip() if row[2] else "",
                "size": str(row[3]).strip() if row[3] else "",
                "cost": float(row[4]) if row[4] is not None else 0,
                "price": float(row[5]) if row[5] is not None else 0,
                "colors": str(row[6]).strip() if row[6] else "",
                "image_url": image_url,
            }
            if not rec["offline_sku"]:
                errors.append(f"第{i}行: 线下SKU 不能为空")
                continue
            records.append(rec)
        except (ValueError, IndexError) as e:
            errors.append(f"第{i}行: 数据格式错误 ({e})")

    if not records:
        raise HTTPException(400, f"没有可导入的有效数据。\n" + "\n".join(errors))

    success, fail, fail_details = db.batch_import(records)
    return JSONResponse({
        "success": success,
        "fail": fail,
        "errors": errors,
        "fail_details": fail_details,
    })


# ── 货款对账 API ─────────────────────────────────────

PAYMENT_IMPORT_HEADERS = ["线下SKU（必填）", "数量（必填）", "成本价（选填）"]


@app.get("/api/payments/import-template")
def payment_import_template(_=Depends(require_auth)):
    """下载货款导入模板"""
    from datetime import date
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "货款导入模板"
    ws.append(list(PAYMENT_IMPORT_HEADERS))
    for col, w in zip("ABC", [18, 14, 14]):
        ws.column_dimensions[col].width = w
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    encoded = quote(f"货款导入模板_{date.today().strftime('%Y%m%d')}.xlsx")
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"}
    )


@app.post("/api/payments/import-items")
def api_import_payment_items(file: UploadFile = File(...), _=Depends(require_auth)):
    """导入货款明细 xlsx，校验 SKU 并返回匹配结果"""
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "仅支持 .xlsx / .xls 文件")

    contents = file.file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(400, "Excel 至少需要表头 + 一行数据")

    header = [str(h).strip() if h else "" for h in rows[0]]
    header_lower = [h.lower() for h in header]

    # 兼容 2 列旧模板 + 3 列新模板（成本价选填）
    valid_headers = [
        ["线下sku", "数量"],                              # 旧版（无括号）
        ["线下sku（必填）", "数量（必填）"],                 # 旧版（有括号）
        ["线下sku（必填）", "数量（必填）", "成本价（选填）"],  # 新版 3 列
    ]
    if header_lower not in valid_headers:
        raise HTTPException(
            400,
            f"表头不匹配，期望: 线下SKU / 数量[ / 成本价（选填）]，实际: {', '.join(header)}"
        )

    raw_items: list[dict] = []
    errors = []
    for i, row in enumerate(rows[1:], start=2):
        if not row or all(v is None for v in row):
            continue
        try:
            offline_sku = str(row[0]).strip() if row[0] is not None else ""
            if not offline_sku:
                errors.append(f"第{i}行: 线下SKU不能为空")
                continue
            if row[1] is None or str(row[1]).strip() == "":
                errors.append(f"第{i}行: 数量不能为空")
                continue
            quantity = float(row[1])
            # 成本价列（可选）：空 → None，有值 → float
            cost_price_override = None
            if len(row) > 2 and row[2] is not None and str(row[2]).strip() != "":
                cost_price_override = float(row[2])
            raw_items.append({
                "offline_sku": offline_sku,
                "quantity": quantity,
                "cost_price_override": cost_price_override,
            })
        except (ValueError, IndexError) as e:
            errors.append(f"第{i}行: 数据格式错误 ({e})")

    if not raw_items:
        return JSONResponse({"items": [], "count": 0, "errors": errors})

    # 批量匹配商品档案（每行独立，不去重合并）
    unique_skus = list({item["offline_sku"] for item in raw_items})
    product_map, not_found = db.match_sku_products(unique_skus)

    # 不存在的 SKU 归入 errors
    for sku in not_found:
        errors.append(f"线下SKU「{sku}」在商品档案中不存在")

    # 按原始顺序逐行组装结果，成本价：文件填值优先，空则取档案成本价
    result_items = []
    for item in raw_items:
        sku = item["offline_sku"]
        if sku in not_found:
            continue
        p = product_map[sku]
        qty = item["quantity"]
        cost_price = item["cost_price_override"] if item["cost_price_override"] is not None else p["cost"]
        result_items.append({
            "offline_sku": sku,
            "quantity": qty,
            "product_name": p["product_name"],
            "cost_price": cost_price,
            "cost_amount": round(cost_price * qty, 2),
        })

    return JSONResponse({
        "items": result_items,
        "count": len(result_items),
        "errors": errors,
    })


@app.get("/api/payments")
def api_list_payments(search: str = "", page: int = 1, page_size: int = 20,
                      _=Depends(require_auth)):
    return db.list_payments(search, page, page_size)


@app.post("/api/payments")
def api_create_payment(data: PaymentCreate, _=Depends(require_auth)):
    items_dict = [item.model_dump() for item in data.items]
    for item in items_dict:
        item["cost_amount"] = round(item["cost_price"] * item["quantity"], 2)
    return db.create_payment(data.name, data.payment_date, items_dict)


@app.get("/api/payments/{payment_id}")
def api_get_payment(payment_id: int, _=Depends(require_auth)):
    p = db.get_payment(payment_id)
    if not p:
        raise HTTPException(404, "货款记录不存在")
    return p


@app.get("/api/payments/{payment_id}/export")
def api_export_payment_items(payment_id: int, _=Depends(require_auth)):
    """导出某笔货款的明细为 Excel"""
    p = db.get_payment(payment_id)
    if not p:
        raise HTTPException(404, "货款记录不存在")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "货款明细"
    ws.append(["商品名称", "线下SKU", "数量", "成本价", "成本额"])
    for item in p.get("items", []):
        ws.append([item["product_name"], item["offline_sku"],
                    item["quantity"], item["cost_price"], item["cost_amount"]])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"{p['name']}_{p['payment_date']}_明细.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    )


@app.put("/api/payments/{payment_id}")
def api_update_payment(payment_id: int, data: PaymentUpdate, _=Depends(require_auth)):
    try:
        items_dict = [item.model_dump() for item in data.items]
        for item in items_dict:
            item["cost_amount"] = round(item["cost_price"] * item["quantity"], 2)
        p = db.update_payment(payment_id, data.name, data.payment_date, items_dict)
    except db.PaymentNotEditableError as e:
        raise HTTPException(400, str(e))
    if not p:
        raise HTTPException(404, "货款记录不存在")
    return p


@app.delete("/api/payments/{payment_id}")
def api_delete_payment(payment_id: int, _=Depends(require_auth)):
    try:
        ok = db.delete_payment(payment_id)
    except db.PaymentNotEditableError as e:
        raise HTTPException(400, str(e))
    if not ok:
        raise HTTPException(404, "货款记录不存在")
    return {"ok": True}


@app.put("/api/payments/{payment_id}/toggle-status")
def api_toggle_payment_status(payment_id: int, _=Depends(require_auth)):
    """撤销结算 → 未结算（保留结算信息）"""
    p = db.toggle_payment_status(payment_id)
    if not p:
        raise HTTPException(404, "货款记录不存在")
    return p


@app.put("/api/payments/{payment_id}/settle")
def api_settle_payment(payment_id: int, data: PaymentSettle, _=Depends(require_auth)):
    """结算货款"""
    p = db.settle_payment(payment_id, data.other_fee, data.settlement_date,
                          data.settlement_remark)
    if not p:
        raise HTTPException(404, "货款记录不存在或已结算")
    return p


@app.post("/api/upload/image")
def api_upload_image(file: UploadFile = File(...), _=Depends(require_auth)):
    """上传结算备注图片"""
    if not file.content_type or not file.content_type.startswith("image/"):
        raise HTTPException(400, "仅支持图片文件")
    ext = os.path.splitext(file.filename or "image.png")[1] or ".png"
    safe_name = f"{secrets.token_hex(16)}{ext}"
    save_path = os.path.join(UPLOADS_DIR, safe_name)
    with open(save_path, "wb") as f:
        f.write(file.file.read())
    return {"url": f"/static/uploads/{safe_name}"}


# ── 页面路由 ─────────────────────────────────────────

STATIC_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static")
UPLOADS_DIR = os.path.join(STATIC_DIR, "uploads")
os.makedirs(STATIC_DIR, exist_ok=True)
os.makedirs(UPLOADS_DIR, exist_ok=True)


@app.get("/")
@app.head("/")
def serve_login():
    """默认入口 → 登录页"""
    return FileResponse(os.path.join(STATIC_DIR, "login.html"))


@app.api_route("/health", methods=["GET", "HEAD"])
def health_check():
    """CloudStudio 健康检查"""
    return {"status": "ok"}


@app.get("/app")
def serve_index():
    """主应用页（登录后跳转）"""
    return FileResponse(os.path.join(STATIC_DIR, "index.html"))


# 挂载 static 目录
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


# ── 启动入口 ─────────────────────────────────────────

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=3888)
